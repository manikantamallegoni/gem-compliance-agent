"""
Document Intelligence module for extracting raw text from submitted files.
Supports PDF (with pytesseract OCR fallback), Word (.docx), and Plain Text (.txt).
"""

from __future__ import annotations

import io
import logging
from fastapi import UploadFile
import docx
import pdfplumber

logger = logging.getLogger("agent.document_intelligence")


async def extract_text_from_file(file: UploadFile) -> str:
    """
    Extract text content from an uploaded PDF, Word document (.docx), or plain text (.txt) file.
    """
    filename = file.filename.lower()
    content = await file.read()

    if filename.endswith(".pdf"):
        return _extract_from_pdf(content)
    elif filename.endswith(".docx"):
        return _extract_from_docx(content)
    elif filename.endswith(".txt"):
        return content.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file format. Please upload a .pdf, .docx, or .txt file.")


def _extract_from_pdf(pdf_bytes: bytes) -> str:
    """
    Extracts text from PDF bytes using pdfplumber, falling back to pytesseract OCR
    if the document is scanned or contains minimal digital text.
    """
    extracted_text = []

    # Standard PDF text extraction via pdfplumber
    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    extracted_text.append(text)
    except Exception as exc:
        logger.warning("pdfplumber text extraction failed: %s", exc)

    full_text = "\n".join(extracted_text)

    # Fallback OCR if text extracted is minimal (scanned PDF case)
    if len(full_text.strip()) < 50:
        logger.info("Minimal native text found (< 50 chars). Triggering pytesseract OCR fallback...")
        full_text = _extract_from_pdf_ocr(pdf_bytes)

    return full_text


def _extract_from_pdf_ocr(pdf_bytes: bytes) -> str:
    """
    Converts PDF pages to PIL images in-memory and runs PyTesseract OCR.
    Requires poppler-utils and tesseract binaries installed on the host system.
    """
    ocr_pages = []
    try:
        from pdf2image import convert_from_bytes
        import pytesseract

        # Convert PDF bytes directly into images
        images = convert_from_bytes(pdf_bytes)
        for i, img in enumerate(images):
            text = pytesseract.image_to_string(img)
            if text.strip():
                ocr_pages.append(text)
            logger.info("OCR completed for PDF page %d", i + 1)
    except Exception as exc:
        logger.warning("OCR processing skipped/failed: %s", exc)

    return "\n".join(ocr_pages)


def _extract_from_docx(docx_bytes: bytes) -> str:
    """
    Extracts paragraphs and structured table content from Word documents (.docx).
    """
    doc = docx.Document(io.BytesIO(docx_bytes))
    text_blocks = []

    # 1. Extract body paragraphs
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_blocks.append(paragraph.text.strip())

    # 2. Extract structured table rows
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_blocks.append(row_text)

    return "\n".join(text_blocks)